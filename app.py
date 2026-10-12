from datetime import date, datetime, timedelta
from contextlib import contextmanager
import math
import os
import random
import sqlite3
import time
import zipfile

from flask import Flask, flash, jsonify, redirect, render_template, request, url_for
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("FLASK_SECRET_KEY", "development-key-change-me")
app.config["DATABASE"] = os.path.join(app.instance_path, "greencampus.db")
DATASET_PATH = os.path.join(os.path.dirname(__file__), "Campus_Sustainability_Dataset_1000 (1).xlsx")

RESOURCES = {
    "electricity": {"label": "Electricity", "unit": "kWh", "threshold": 1000},
    "water": {"label": "Water", "unit": "L", "threshold": 10000},
    "waste": {"label": "Waste", "unit": "kg", "threshold": 100},
}
BUILDINGS = [
    "Main Block",
    "Science Block",
    "Library",
    "Hostel",
    "Cafeteria",
    "Sports Complex",
]
AUTOPILOT_INTERVAL_SECONDS = 30
SIMULATED_NOTE_PREFIX = "[SIMULATED]"
RESOURCE_ACTIONS = {
    "electricity": "Inspect lights, HVAC, and equipment in unused rooms; switch off what is not needed.",
    "water": "Check nearby taps, toilets, and pipes for leaks and report any issue.",
    "waste": "Check collection points and improve sorting, recycling, or composting signage.",
}
DATASET_COLUMNS = (
    "Record_ID",
    "Date",
    "Building",
    "Day_Type",
    "Occupancy_People",
    "Electricity_Consumption_kWh",
    "Water_Consumption_Liters",
    "Waste_Generated_kg",
    "Renewable_Energy_Share_pct",
    "Renewable_Energy_kWh",
    "Estimated_CO2_kg",
    "Energy_Status",
    "Water_Status",
    "Suggested_Action",
)
DATASET_RESOURCE_COLUMNS = {
    "electricity": "Electricity_Consumption_kWh",
    "water": "Water_Consumption_Liters",
    "waste": "Waste_Generated_kg",
}


def parse_iso_date(value):
    parsed = datetime.strptime(value, "%Y-%m-%d").date()
    if parsed.isoformat() != value:
        raise ValueError("Date must use YYYY-MM-DD format.")
    return parsed


@contextmanager
def get_db():
    os.makedirs(app.instance_path, exist_ok=True)
    connection = sqlite3.connect(app.config["DATABASE"])
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
    except Exception:
        connection.rollback()
        raise
    else:
        connection.commit()
    finally:
        connection.close()


def init_db():
    with get_db() as db:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS resource_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                building TEXT NOT NULL,
                resource_type TEXT NOT NULL,
                quantity REAL NOT NULL CHECK(quantity > 0),
                record_date TEXT NOT NULL,
                notes TEXT NOT NULL DEFAULT '',
                is_demo INTEGER NOT NULL DEFAULT 0,
                demo_key TEXT UNIQUE,
                source_type TEXT NOT NULL DEFAULT 'MEASURED',
                source_record_id TEXT,
                day_type TEXT,
                occupancy_people REAL,
                renewable_energy_share_pct REAL,
                renewable_energy_kwh REAL,
                estimated_co2_kg REAL,
                energy_status TEXT,
                water_status TEXT,
                suggested_action TEXT
            )
            """
        )
        existing_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(resource_records)")
        }
        additions = {
            "source_type": "TEXT NOT NULL DEFAULT 'MEASURED'",
            "source_record_id": "TEXT",
            "day_type": "TEXT",
            "occupancy_people": "REAL",
            "renewable_energy_share_pct": "REAL",
            "renewable_energy_kwh": "REAL",
            "estimated_co2_kg": "REAL",
            "energy_status": "TEXT",
            "water_status": "TEXT",
            "suggested_action": "TEXT",
        }
        for column, definition in additions.items():
            if column not in existing_columns:
                db.execute(
                    f"ALTER TABLE resource_records ADD COLUMN {column} {definition}"
                )
        db.execute(
            """UPDATE resource_records SET source_type = 'DEMO'
               WHERE is_demo = 1 AND source_type = 'MEASURED'"""
        )
        db.execute(
            """UPDATE resource_records SET source_type = 'SIMULATED'
               WHERE substr(notes, 1, 11) = '[SIMULATED]' AND source_type = 'MEASURED'"""
        )
        db.execute(
            """CREATE UNIQUE INDEX IF NOT EXISTS uq_dataset_record_resource
               ON resource_records (source_type, source_record_id, resource_type)
               WHERE source_type = 'DATASET' AND source_record_id IS NOT NULL"""
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS autopilot_state (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                is_running INTEGER NOT NULL DEFAULT 0,
                building TEXT NOT NULL,
                next_batch_at REAL NOT NULL DEFAULT 0
            )
            """
        )
        db.execute(
            """INSERT OR IGNORE INTO autopilot_state (id, is_running, building, next_batch_at)
               VALUES (1, 0, ?, 0)""",
            (BUILDINGS[0],),
        )


def optional_dataset_number(value, field_name):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{field_name} must be numeric.")
    number = float(value)
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"{field_name} must be a finite, non-negative number.")
    return number


def dataset_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return parse_iso_date(value.strip()).isoformat()
    raise ValueError("Date is missing or invalid.")


def import_dataset_file(path=None):
    workbook_path = path or DATASET_PATH
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Dataset" not in workbook.sheetnames:
            raise ValueError("The workbook must contain a sheet named Dataset.")
        sheet = workbook["Dataset"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("The Dataset worksheet is empty.")
        header_indexes = {
            str(header).strip(): index
            for index, header in enumerate(headers)
            if header is not None
        }
        missing_headers = [name for name in DATASET_COLUMNS if name not in header_indexes]
        if missing_headers:
            raise ValueError(
                "The Dataset worksheet is missing columns: "
                + ", ".join(missing_headers)
            )

        imported = 0
        skipped = 0
        errors = []
        with get_db() as db:
            for sheet_row_number, values in enumerate(rows, start=2):
                if not values or not any(value is not None for value in values):
                    continue
                record = {
                    name: values[index] if index < len(values) else None
                    for name, index in header_indexes.items()
                    if name in DATASET_COLUMNS
                }
                try:
                    raw_id = record["Record_ID"]
                    record_id = str(raw_id).strip() if raw_id is not None else ""
                    building = record["Building"]
                    building = building.strip() if isinstance(building, str) else ""
                    if not record_id or len(record_id) > 100:
                        raise ValueError("Record_ID is missing or invalid.")
                    if not building or len(building) > 200:
                        raise ValueError("Building is missing or invalid.")
                    record_date = dataset_date(record["Date"])
                    occupancy = optional_dataset_number(
                        record["Occupancy_People"], "Occupancy_People"
                    )
                    if occupancy is not None and not occupancy.is_integer():
                        raise ValueError("Occupancy_People must be a whole number.")
                    renewable_share = optional_dataset_number(
                        record["Renewable_Energy_Share_pct"],
                        "Renewable_Energy_Share_pct",
                    )
                    if renewable_share is not None and renewable_share > 100:
                        raise ValueError("Renewable_Energy_Share_pct cannot exceed 100.")
                    renewable_kwh = optional_dataset_number(
                        record["Renewable_Energy_kWh"], "Renewable_Energy_kWh"
                    )
                    estimated_co2 = optional_dataset_number(
                        record["Estimated_CO2_kg"], "Estimated_CO2_kg"
                    )
                    metadata = {
                        "day_type": record["Day_Type"],
                        "occupancy_people": occupancy,
                        "renewable_energy_share_pct": renewable_share,
                        "renewable_energy_kwh": renewable_kwh,
                        "estimated_co2_kg": estimated_co2,
                        "energy_status": record["Energy_Status"],
                        "water_status": record["Water_Status"],
                        "suggested_action": record["Suggested_Action"],
                    }
                    for key, value in metadata.items():
                        if isinstance(value, str):
                            value = value.strip() or None
                            metadata[key] = value
                        elif value is not None and key in (
                            "day_type",
                            "energy_status",
                            "water_status",
                            "suggested_action",
                        ):
                            raise ValueError(f"{key} must be text.")
                    readings = []
                    for resource_type, column in DATASET_RESOURCE_COLUMNS.items():
                        quantity = optional_dataset_number(record[column], column)
                        if quantity is None or quantity <= 0:
                            raise ValueError(f"{column} must be greater than zero.")
                        readings.append((resource_type, quantity))
                except (TypeError, ValueError, OverflowError) as error:
                    skipped += 1
                    if len(errors) < 5:
                        errors.append(f"Row {sheet_row_number}: {error}")
                    continue

                inserted_for_row = 0
                for resource_type, quantity in readings:
                    cursor = db.execute(
                        """INSERT OR IGNORE INTO resource_records
                           (building, resource_type, quantity, record_date, notes,
                            is_demo, source_type, source_record_id, day_type,
                            occupancy_people, renewable_energy_share_pct,
                            renewable_energy_kwh, estimated_co2_kg, energy_status,
                            water_status, suggested_action)
                           VALUES (?, ?, ?, ?, ?, 0, 'DATASET', ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (
                            building,
                            resource_type,
                            quantity,
                            record_date,
                            f"[DATASET] Imported workbook row; Record_ID={record_id}",
                            record_id,
                            metadata["day_type"],
                            metadata["occupancy_people"],
                            metadata["renewable_energy_share_pct"],
                            metadata["renewable_energy_kwh"],
                            metadata["estimated_co2_kg"],
                            metadata["energy_status"],
                            metadata["water_status"],
                            metadata["suggested_action"],
                        ),
                    )
                    inserted_for_row += cursor.rowcount
                if inserted_for_row:
                    imported += 1
                else:
                    skipped += 1
        return {"imported": imported, "skipped": skipped, "errors": errors}
    finally:
        workbook.close()


def get_autopilot_status(now=None):
    with get_db() as db:
        state = db.execute(
            "SELECT is_running, building, next_batch_at FROM autopilot_state WHERE id = 1"
        ).fetchone()
    current_time = time.time() if now is None else now
    return {
        "running": bool(state["is_running"]),
        "building": state["building"],
        "seconds_until_batch": (
            max(0, math.ceil(state["next_batch_at"] - current_time))
            if state["is_running"]
            else None
        ),
    }


def make_simulated_readings(building, today):
    readings = []
    for resource_type, normal_range, high_range in (
        ("electricity", (35, 160), (1100, 1500)),
        ("water", (300, 2500), (11000, 14000)),
        ("waste", (5, 35), (105, 150)),
    ):
        # Occasional spikes make the threshold-alert demo visible without making every batch abnormal.
        quantity_range = high_range if random.random() < 0.15 else normal_range
        quantity = round(random.uniform(*quantity_range), 1)
        readings.append(
            (
                building,
                resource_type,
                quantity,
                today.isoformat(),
                f"{SIMULATED_NOTE_PREFIX} AutoPilot sample; not sensor data.",
            )
        )
    return readings


def get_dashboard_snapshot():
    totals = get_totals()
    with get_db() as db:
        daily = db.execute(
            """SELECT record_date, resource_type, SUM(quantity) AS total
               FROM resource_records GROUP BY record_date, resource_type
               ORDER BY record_date"""
        ).fetchall()
        records = db.execute(
            """SELECT id, building, resource_type, quantity, record_date, notes, is_demo,
                      source_type, source_record_id
               FROM resource_records ORDER BY record_date DESC, id DESC LIMIT 6"""
        ).fetchall()
        score_rows = db.execute(
            "SELECT resource_type, quantity FROM resource_records"
        ).fetchall()

    labels = sorted({row["record_date"] for row in daily})
    series = {kind: {label: 0 for label in labels} for kind in RESOURCES}
    for row in daily:
        if row["resource_type"] in series:
            series[row["resource_type"]][row["record_date"]] = row["total"]
    score = (
        100
        if not score_rows
        else max(
            0,
            round(
                100
                - sum(
                    row["quantity"] > RESOURCES[row["resource_type"]]["threshold"]
                    for row in score_rows
                )
                / len(score_rows)
                * 100
            ),
        )
    )
    return {
        "totals": totals,
        "score": score,
        "recommendations": [
            {"title": title, "suggestion": suggestion}
            for title, suggestion in get_recommendations(totals)
        ],
        "chart_data": {
            "labels": labels,
            "series": {
                kind: [series[kind][label] for label in labels] for kind in RESOURCES
            },
        },
        "records": [
            {
                "id": row["id"],
                "building": row["building"],
                "resource_type": row["resource_type"],
                "quantity": row["quantity"],
                "record_date": row["record_date"],
                "notes": row["notes"],
                "is_demo": bool(row["is_demo"]),
                "source_type": row["source_type"],
                "source_record_id": row["source_record_id"],
                "is_simulated": row["source_type"] == "SIMULATED",
            }
            for row in records
        ],
    }


def process_autopilot_tick(now=None):
    current_time = time.time() if now is None else now
    with get_db() as db:
        db.execute("BEGIN IMMEDIATE")
        state = db.execute(
            "SELECT is_running, building, next_batch_at FROM autopilot_state WHERE id = 1"
        ).fetchone()
        generated_ids = []
        if state["is_running"] and current_time >= state["next_batch_at"]:
            readings = make_simulated_readings(state["building"], date.today())
            for reading in readings:
                cursor = db.execute(
                    """INSERT INTO resource_records
                       (building, resource_type, quantity, record_date, notes, is_demo, source_type)
                       VALUES (?, ?, ?, ?, ?, 0, 'SIMULATED')""",
                    reading,
                )
                generated_ids.append(cursor.lastrowid)
            db.execute(
                "UPDATE autopilot_state SET next_batch_at = ? WHERE id = 1",
                (current_time + AUTOPILOT_INTERVAL_SECONDS,),
            )
        state = db.execute(
            "SELECT is_running, building, next_batch_at FROM autopilot_state WHERE id = 1"
        ).fetchone()
        generated = []
        if generated_ids:
            placeholders = ",".join("?" for _ in generated_ids)
            generated = db.execute(
                f"SELECT id, building, resource_type, quantity, record_date FROM resource_records WHERE id IN ({placeholders}) ORDER BY id",
                generated_ids,
            ).fetchall()

    alerts = []
    for row in generated:
        resource = RESOURCES[row["resource_type"]]
        if row["quantity"] > resource["threshold"]:
            alerts.append(
                {
                    "id": row["id"],
                    "building": row["building"],
                    "resource": resource["label"],
                    "quantity": row["quantity"],
                    "unit": resource["unit"],
                    "threshold": resource["threshold"],
                    "action": RESOURCE_ACTIONS[row["resource_type"]],
                }
            )
    return {
        "status": {
            "running": bool(state["is_running"]),
            "building": state["building"],
            "seconds_until_batch": (
                max(0, math.ceil(state["next_batch_at"] - current_time))
                if state["is_running"]
                else None
            ),
        },
        "generated": [
            {
                "id": row["id"],
                "resource_type": row["resource_type"],
                "quantity": row["quantity"],
            }
            for row in generated
        ],
        "alerts": alerts,
        "snapshot": get_dashboard_snapshot(),
    }


def get_totals(start_date=None, end_date=None, source_type=None):
    totals = {kind: 0.0 for kind in RESOURCES}
    clauses = []
    params = []
    if start_date:
        clauses.append("record_date >= ?")
        params.append(start_date)
    if end_date:
        clauses.append("record_date <= ?")
        params.append(end_date)
    if source_type:
        clauses.append("source_type = ?")
        params.append(source_type)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    with get_db() as db:
        rows = db.execute(
            f"""SELECT resource_type, SUM(quantity) AS total
                FROM resource_records {where} GROUP BY resource_type""",
            params,
        ).fetchall()
    for row in rows:
        totals[row["resource_type"]] = row["total"] or 0.0
    return totals


def get_alerts():
    with get_db() as db:
        rows = db.execute(
            "SELECT * FROM resource_records ORDER BY record_date DESC, id DESC"
        ).fetchall()
    alerts = []
    for row in rows:
        threshold = RESOURCES[row["resource_type"]]["threshold"]
        if row["quantity"] > threshold:
            alerts.append(
                {
                    "record": row,
                    "threshold": threshold,
                    "resource": RESOURCES[row["resource_type"]],
                    "action": (
                        row["suggested_action"]
                        if row["source_type"] == "DATASET" and row["suggested_action"]
                        else RESOURCE_ACTIONS[row["resource_type"]]
                    ),
                }
            )
    return alerts


@app.context_processor
def inject_today_display():
    return {"today_display": date.today().strftime("%A, %B %d, %Y")}


def get_recommendations(totals):
    recommendations = []
    if totals["electricity"] > 0:
        recommendations.append(
            ("Electricity", "Switch off lights, projectors, and air conditioning in unused rooms.")
        )
    if totals["water"] > 0:
        recommendations.append(
            ("Water", "Check taps and washrooms for leaks, and report repairs promptly.")
        )
    if totals["waste"] > 0:
        recommendations.append(
            ("Waste", "Place clearly labeled recycling and compost bins near high-traffic areas.")
        )
    if not recommendations:
        recommendations.append(
            ("Get started", "Add a resource record to receive practical, resource-specific actions.")
        )
    return recommendations


@app.route("/")
def dashboard():
    totals = get_totals()
    with get_db() as db:
        records = db.execute(
            "SELECT * FROM resource_records ORDER BY record_date DESC, id DESC LIMIT 6"
        ).fetchall()
        sample_count = db.execute(
            "SELECT COUNT(*) FROM resource_records WHERE is_demo = 1"
        ).fetchone()[0]
        dataset_count = db.execute(
            "SELECT COUNT(*) FROM resource_records WHERE source_type = 'DATASET'"
        ).fetchone()[0]
        record_count = db.execute("SELECT COUNT(*) FROM resource_records").fetchone()[0]
        daily = db.execute(
            """SELECT record_date, resource_type, SUM(quantity) AS total
               FROM resource_records GROUP BY record_date, resource_type
               ORDER BY record_date"""
        ).fetchall()

    score = 100 if record_count == 0 else max(0, round(100 - (len(get_alerts()) / record_count * 100)))
    chart_data = {
        "labels": sorted({row["record_date"] for row in daily}),
        "series": {
            kind: [
                next(
                    (
                        row["total"]
                        for row in daily
                        if row["record_date"] == record_date and row["resource_type"] == kind
                    ),
                    0,
                )
                for record_date in sorted({row["record_date"] for row in daily})
            ]
            for kind in RESOURCES
        },
    }
    return render_template(
        "dashboard.html",
        resources=RESOURCES,
        buildings=BUILDINGS,
        totals=totals,
        records=records,
        alerts=get_alerts(),
        recommendations=get_recommendations(totals),
        score=score,
        chart_data=chart_data,
        sample_count=sample_count,
        dataset_count=dataset_count,
        today=date.today().isoformat(),
        autopilot_status=get_autopilot_status(),
    )


@app.post("/dataset/import")
def import_dataset():
    try:
        result = import_dataset_file()
    except FileNotFoundError:
        flash(
            "Dataset workbook not found. Place Campus_Sustainability_Dataset_1000 (1).xlsx beside app.py.",
            "error",
        )
        return redirect(url_for("dashboard"))
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as error:
        flash(f"Dataset import failed: {error}", "error")
        return redirect(url_for("dashboard"))

    flash(
        "Dataset import complete: "
        f"{result['imported']} source rows imported, "
        f"{result['skipped']} rows skipped (invalid rows or already imported).",
        "success",
    )
    for error in result["errors"][:3]:
        flash(error, "error")
    return redirect(url_for("dashboard"))


@app.get("/api/autopilot/status")
def autopilot_status():
    return jsonify(get_autopilot_status())


@app.post("/api/autopilot/start")
def start_autopilot():
    if not request.is_json:
        return jsonify(error="Send the selected building as JSON."), 400
    payload = request.get_json(silent=True)
    building = payload.get("building") if isinstance(payload, dict) else None
    if building not in BUILDINGS:
        return jsonify(error="Choose a valid campus building."), 400

    current_time = time.time()
    with get_db() as db:
        db.execute("BEGIN IMMEDIATE")
        state = db.execute(
            "SELECT is_running, next_batch_at FROM autopilot_state WHERE id = 1"
        ).fetchone()
        next_batch_at = (
            state["next_batch_at"]
            if state["is_running"]
            else current_time + AUTOPILOT_INTERVAL_SECONDS
        )
        db.execute(
            """UPDATE autopilot_state
               SET is_running = 1, building = ?, next_batch_at = ?
               WHERE id = 1""",
            (building, next_batch_at),
        )
    return jsonify(get_autopilot_status())


@app.post("/api/autopilot/stop")
def stop_autopilot():
    with get_db() as db:
        db.execute("BEGIN IMMEDIATE")
        db.execute(
            "UPDATE autopilot_state SET is_running = 0, next_batch_at = 0 WHERE id = 1"
        )
    return jsonify(get_autopilot_status())


@app.post("/api/autopilot/tick")
def autopilot_tick():
    return jsonify(process_autopilot_tick())


@app.post("/records")
def add_record():
    building = request.form.get("building", "").strip()
    resource_type = request.form.get("resource_type", "").strip()
    quantity_text = request.form.get("quantity", "").strip()
    record_date = request.form.get("record_date", "").strip()
    notes = request.form.get("notes", "").strip()

    errors = []
    if building not in BUILDINGS:
        errors.append("Choose a valid campus building.")
    if resource_type not in RESOURCES:
        errors.append("Choose a valid resource type.")
    try:
        quantity = float(quantity_text)
        if not (0 < quantity < 1_000_000_000):
            raise ValueError
    except ValueError:
        errors.append("Quantity must be a positive number below 1,000,000,000.")
        quantity = 0
    try:
        parsed_date = parse_iso_date(record_date)
        if parsed_date > date.today():
            errors.append("Record date cannot be in the future.")
    except ValueError:
        errors.append("Enter a valid record date.")

    if len(notes) > 500:
        errors.append("Notes must be 500 characters or fewer.")
    if errors:
        for error in errors:
            flash(error, "error")
        return redirect(url_for("dashboard") + "#add-record")

    with get_db() as db:
        db.execute(
            """INSERT INTO resource_records
               (building, resource_type, quantity, record_date, notes)
               VALUES (?, ?, ?, ?, ?)""",
            (building, resource_type, quantity, record_date, notes),
        )
    flash("Resource record saved.", "success")
    return redirect(url_for("dashboard"))


@app.route("/history")
def history():
    resource_type = request.args.get("resource_type", "").strip()
    source_type = request.args.get("source_type", "").strip().upper()
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    errors = []
    clauses = []
    params = []
    parsed_start = None
    parsed_end = None
    if resource_type in RESOURCES:
        clauses.append("resource_type = ?")
        params.append(resource_type)
    elif resource_type:
        errors.append("Choose a valid resource type.")
    valid_sources = {"MEASURED", "DEMO", "SIMULATED", "DATASET"}
    if source_type:
        if source_type not in valid_sources:
            errors.append("Choose a valid record source.")
        else:
            clauses.append("source_type = ?")
            params.append(source_type)
    if start_date:
        try:
            parsed_start = parse_iso_date(start_date)
            clauses.append("record_date >= ?")
            params.append(parsed_start.isoformat())
        except ValueError:
            errors.append("Enter a valid start date in YYYY-MM-DD format.")
    if end_date:
        try:
            parsed_end = parse_iso_date(end_date)
            clauses.append("record_date <= ?")
            params.append(parsed_end.isoformat())
        except ValueError:
            errors.append("Enter a valid end date in YYYY-MM-DD format.")
    if parsed_start and parsed_end and parsed_start > parsed_end:
        errors.append("The start date must be on or before the end date.")
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    if errors:
        records = []
        for error in errors:
            flash(error, "error")
    else:
        with get_db() as db:
            records = db.execute(
                f"SELECT * FROM resource_records {where} ORDER BY record_date DESC, id DESC",
                params,
            ).fetchall()
    return render_template(
        "history.html",
        records=records,
        resources=RESOURCES,
        filters={
            "resource_type": resource_type,
            "source_type": source_type,
            "start_date": start_date,
            "end_date": end_date,
        },
    )


@app.post("/records/<int:record_id>/delete")
def delete_record(record_id):
    with get_db() as db:
        result = db.execute("DELETE FROM resource_records WHERE id = ?", (record_id,))
    if result.rowcount:
        flash("Resource record deleted.", "success")
    else:
        flash("That resource record was not found.", "error")
    return redirect(url_for("history"))


@app.post("/demo-data")
def load_demo_data():
    today = date.today()
    demo_rows = [
        ("Main Block", "electricity", 420, (today - timedelta(days=6)).isoformat(), "Sample demo reading", "demo-elec-1"),
        ("Library", "electricity", 510, (today - timedelta(days=3)).isoformat(), "Sample demo reading", "demo-elec-2"),
        ("Science Block", "water", 2800, (today - timedelta(days=5)).isoformat(), "Sample demo reading", "demo-water-1"),
        ("Hostel", "water", 12500, (today - timedelta(days=2)).isoformat(), "Sample demo reading: above the water alert threshold", "demo-water-2"),
        ("Cafeteria", "waste", 74, (today - timedelta(days=1)).isoformat(), "Sample demo reading", "demo-waste-1"),
    ]
    with get_db() as db:
        changes_before = db.total_changes
        db.executemany(
            """INSERT OR IGNORE INTO resource_records
               (building, resource_type, quantity, record_date, notes, is_demo, demo_key, source_type)
               VALUES (?, ?, ?, ?, ?, 1, ?, 'DEMO')""",
            demo_rows,
        )
        added = db.total_changes - changes_before
    if added:
        flash(f"Loaded {added} clearly labeled sample records.", "success")
    else:
        flash("Demo data is already loaded; no duplicate records were added.", "success")
    return redirect(url_for("dashboard"))


@app.route("/impact", methods=["GET", "POST"])
def impact():
    today = date.today()
    default_end = today.isoformat()
    default_start = (today - timedelta(days=13)).isoformat()
    period_one_start = request.values.get("period_one_start", default_start)
    period_one_end = request.values.get("period_one_end", (today - timedelta(days=7)).isoformat())
    period_two_start = request.values.get("period_two_start", (today - timedelta(days=6)).isoformat())
    period_two_end = request.values.get("period_two_end", default_end)
    source_type = request.values.get("source_type", "").strip().upper()
    results = None
    if request.method == "POST":
        try:
            if source_type and source_type not in {"MEASURED", "DEMO", "SIMULATED", "DATASET"}:
                raise ValueError
            ranges = [
                (parse_iso_date(period_one_start), parse_iso_date(period_one_end)),
                (parse_iso_date(period_two_start), parse_iso_date(period_two_end)),
            ]
            if any(start > end for start, end in ranges):
                raise ValueError
            first = get_totals(period_one_start, period_one_end, source_type or None)
            second = get_totals(period_two_start, period_two_end, source_type or None)
            results = {}
            for kind in RESOURCES:
                previous, current = first[kind], second[kind]
                change = None if previous == 0 else ((current - previous) / previous) * 100
                results[kind] = {"previous": previous, "current": current, "change": change}
        except ValueError:
            flash(
                "Choose a valid record source and date ranges with each start date on or before its end date.",
                "error",
            )
    return render_template(
        "impact.html",
        resources=RESOURCES,
        results=results,
        periods=(period_one_start, period_one_end, period_two_start, period_two_end),
        source_type=source_type,
    )


with app.app_context():
    init_db()


if __name__ == "__main__":
    app.run(debug=True)
